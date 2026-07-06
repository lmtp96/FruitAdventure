from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(r"D:\Game Project\FruitAdventure\output\documents\Mau_Bia_Bao_Cao_Fruit_Adventure.docx")


def set_run_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def paragraph(text="", size=13, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic)
    return p


def add_page_border(section):
    sect_pr = section._sectPr
    pg_borders = OxmlElement("w:pgBorders")
    pg_borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "double")
        border.set(qn("w:sz"), "12")
        border.set(qn("w:space"), "24")
        border.set(qn("w:color"), "000000")
        pg_borders.append(border)
    sect_pr.append(pg_borders)


def add_logo_placeholder():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Cm(5.2)
    p.paragraph_format.right_indent = Cm(5.2)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15

    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:space"), "12")
        border.set(qn("w:color"), "B7B7B7")
        borders.append(border)
    p_pr.append(borders)

    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "420")
    spacing.set(qn("w:after"), "420")
    p_pr.append(spacing)

    r = p.add_run("[CHÈN LOGO HỌC VIỆN TẠI ĐÂY]")
    set_run_font(r, size=13, bold=True, color=(80, 80, 80))


def add_info_line(label, value, after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(3.0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.tab_stops.add_tab_stop(Cm(8.4))
    label_run = p.add_run(label)
    set_run_font(label_run, size=14, bold=False)
    value_run = p.add_run(f"\t{value}")
    set_run_font(value_run, size=14, bold=True)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.4)
section.right_margin = Cm(2.4)
section.header_distance = Cm(1.25)
section.footer_distance = Cm(1.25)
add_page_border(section)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(13)

for style_name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 13)):
    style = styles[style_name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15

paragraph("HỌC VIỆN CÔNG NGHỆ THÔNG TIN & THIẾT KẾ VTC", size=16, bold=True, after=0)
paragraph("ACADEMY", size=16, bold=True, after=22)

add_logo_placeholder()

paragraph("CAPSTONE PROJECT", size=18, bold=True, before=30, after=44)
paragraph("LẬP TRÌNH GAME 2D PLATFORMER", size=20, bold=True, after=4)
paragraph("FRUIT ADVENTURE", size=22, bold=True, after=28)

add_info_line("Ngành:", "CÔNG NGHỆ THÔNG TIN", after=4)
add_info_line("Chuyên ngành:", "LẬP TRÌNH GAME", after=26)
add_info_line("Giảng viên hướng dẫn:", "[Tên giảng viên]", after=4)
add_info_line("Học viên thực hiện:", "[Tên học viên]", after=4)
add_info_line("MSHV:", "[Mã số học viên]        Lớp: [Tên lớp]", after=4)

paragraph("TP. Hồ Chí Minh, 2026", size=15, before=54, after=0)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
